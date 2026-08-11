from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (
    Attendance,
    Camera,
    EventType,
    ExamSession,
    Identity,
    RecognitionEvent,
    SessionMember,
    User,
)
from .storage import blobs


NAVY = "16202B"
TEAL = "24C7A5"
RED = "E45D5D"
MUTED = "E9EDF2"


def _local_text(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")


def _header(sheet, row: int, values: list[str], fill: str = NAVY) -> None:  # type: ignore[no-untyped-def]
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24


def _autosize(sheet, widths: dict[int, int] | None = None) -> None:  # type: ignore[no-untyped-def]
    selected = widths or {}
    for column_cells in sheet.columns:
        column = column_cells[0].column
        if column in selected:
            width = selected[column]
        else:
            width = min(45, max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width


def _add_event_sheet(
    workbook: Workbook,
    db: Session,
    title: str,
    events: list[RecognitionEvent],
    *,
    fill: str,
) -> None:
    sheet = workbook.create_sheet(title)
    _header(
        sheet,
        1,
        [
            "Captured face",
            "Time",
            "Type",
            "Detected identity",
            "Candidate No.",
            "Camera",
            "Confidence",
            "Resolution",
            "Resolved by",
            "Remarks",
            "Telegram",
        ],
        fill=fill,
    )
    for row_no, event in enumerate(events, start=2):
        camera = db.get(Camera, event.camera_id)
        resolved = db.get(Identity, event.resolved_identity_id) if event.resolved_identity_id else None
        resolver = db.get(User, event.resolved_by_user_id) if event.resolved_by_user_id else None
        sheet.append(
            [
                "",
                _local_text(event.captured_at),
                event.event_type.replace("_", " ").title(),
                event.identity_name_snapshot or "Unknown",
                event.candidate_no_snapshot or "",
                camera.name if camera else str(event.camera_id),
                round(event.confidence, 4) if event.confidence is not None else "",
                f"Manual match: {resolved.display_name}" if resolved else "Unresolved",
                resolver.display_name if resolver else "",
                event.resolution_note or "",
                event.telegram_status.replace("_", " ").title(),
            ]
        )
        sheet.row_dimensions[row_no].height = 72
        try:
            face_bytes = blobs.get(event.snapshot_blob_key)
            image = Image.open(io.BytesIO(face_bytes))
            image.thumbnail((82, 82))
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            buffer.seek(0)
            worksheet_image = WorksheetImage(buffer)
            worksheet_image.width = image.width
            worksheet_image.height = image.height
            sheet.add_image(worksheet_image, f"A{row_no}")
        except Exception:
            sheet.cell(row_no, 1, "Image unavailable")
        for cell in sheet[row_no]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    _autosize(sheet, {1: 15, 2: 24, 4: 26, 8: 32, 10: 42})


def build_session_report(db: Session, session_id: int) -> tuple[str, bytes]:
    session = db.get(ExamSession, session_id)
    if not session:
        raise ValueError("Exam session not found")

    member_rows = db.execute(
        select(SessionMember, Identity)
        .join(Identity, Identity.id == SessionMember.identity_id)
        .where(SessionMember.session_id == session_id)
        .order_by(Identity.display_name)
    ).all()
    attendance_rows = db.scalars(
        select(Attendance).where(Attendance.session_id == session_id)
    ).all()
    attendance_by_identity = {row.identity_id: row for row in attendance_rows}
    events = db.scalars(
        select(RecognitionEvent)
        .where(RecognitionEvent.session_id == session_id)
        .order_by(RecognitionEvent.captured_at)
    ).all()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "FAST-FR EXAM SESSION REPORT"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary.merge_cells("A1:D1")
    fields = [
        ("Session", session.name),
        ("Code", session.code),
        ("Starts", _local_text(session.starts_at)),
        ("Ends", _local_text(session.ends_at)),
        ("Status", session.status.title()),
        ("Generated", _local_text(datetime.now(timezone.utc))),
        ("Roster", len(member_rows)),
        ("Present", len(attendance_rows)),
        ("Absent", max(0, len(member_rows) - len(attendance_rows))),
        ("Unknown", sum(event.event_type == EventType.UNKNOWN.value for event in events)),
        (
            "Not registered",
            sum(event.event_type == EventType.NOT_REGISTERED.value for event in events),
        ),
        ("Blacklist alerts", sum(event.event_type == EventType.BLACKLIST.value for event in events)),
    ]
    for row_no, (label, value) in enumerate(fields, start=3):
        summary.cell(row_no, 1, label).font = Font(bold=True, color=NAVY)
        summary.cell(row_no, 2, value)
    summary.column_dimensions["A"].width = 23
    summary.column_dimensions["B"].width = 42
    summary.freeze_panes = "A3"

    attendance_sheet = workbook.create_sheet("Attendance")
    headers = [
        "Candidate No.",
        "Name",
        "Status",
        "Source",
        "First seen",
        "Last seen",
        "Camera",
        "Confidence",
        "Operator",
        "Remarks",
    ]
    _header(attendance_sheet, 1, headers)
    for row_no, (_membership, identity) in enumerate(member_rows, start=2):
        attendance = attendance_by_identity.get(identity.id)
        camera = db.get(Camera, attendance.camera_id) if attendance and attendance.camera_id else None
        operator = (
            db.get(User, attendance.operator_user_id)
            if attendance and attendance.operator_user_id
            else None
        )
        values = [
            identity.candidate_no or "",
            identity.display_name,
            "Present" if attendance else "Absent",
            attendance.source.replace("_", " ").title() if attendance else "",
            _local_text(attendance.first_seen_at) if attendance else "",
            _local_text(attendance.last_seen_at) if attendance else "",
            camera.name if camera else "",
            round(attendance.confidence, 4) if attendance and attendance.confidence else "",
            operator.display_name if operator else "",
            attendance.remarks if attendance else "",
        ]
        attendance_sheet.append(values)
        fill = "E5F7F2" if attendance else "F5F6F8"
        for cell in attendance_sheet[row_no]:
            cell.fill = PatternFill("solid", fgColor=fill)
    attendance_sheet.freeze_panes = "A2"
    attendance_sheet.auto_filter.ref = attendance_sheet.dimensions
    _autosize(attendance_sheet, {2: 28, 5: 24, 6: 24, 10: 40})

    _add_event_sheet(
        workbook,
        db,
        "Unknowns",
        [event for event in events if event.event_type in {
            EventType.UNKNOWN.value,
            EventType.NOT_REGISTERED.value,
        }],
        fill=RED,
    )
    _add_event_sheet(
        workbook,
        db,
        "Blacklist",
        [event for event in events if event.event_type == EventType.BLACKLIST.value],
        fill="7A1018",
    )

    safe_code = "".join(character for character in session.code if character.isalnum() or character in "-_")
    filename = f"FastFR-{safe_code or session.id}-attendance.xlsx"
    output = io.BytesIO()
    workbook.save(output)
    return filename, output.getvalue()
