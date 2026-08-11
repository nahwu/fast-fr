from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from fastfr.db import SessionLocal
from fastfr.main import app
from fastfr.models import (
    Attendance,
    Camera,
    EventType,
    RecognitionEvent,
)
from fastfr.realtime import event_hub
from fastfr.security import crypto
from fastfr.storage import blobs


def jpeg_bytes(color: str = "#d7b18a") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (96, 112), color).save(output, "JPEG")
    return output.getvalue()


def test_end_to_end_admin_operator_reporting_and_websocket() -> None:
    with TestClient(app) as client:
        assert client.get("/api/setup/status").json() == {"required": True}
        setup = client.post(
            "/api/setup/admin",
            json={
                "username": "admin",
                "displayName": "System Administrator",
                "password": "StrongPassword123!",
            },
        )
        assert setup.status_code == 201
        csrf = setup.json()["csrfToken"]
        headers = {"X-CSRF-Token": csrf}

        assert client.post(
            "/api/identities",
            json={"displayName": "Missing CSRF", "category": "whitelist"},
        ).status_code == 403
        candidate = client.post(
            "/api/identities",
            headers=headers,
            json={"displayName": "Jane Candidate", "candidateNo": "EX-001", "category": "whitelist"},
        ).json()
        blocked = client.post(
            "/api/identities",
            headers=headers,
            json={"displayName": "Blocked Person", "category": "blacklist"},
        ).json()
        session = client.post(
            "/api/sessions",
            headers=headers,
            json={
                "name": "Engineering Entrance Exam",
                "code": "ENG-2026-A",
                "startsAt": "2026-08-12T01:00:00Z",
                "endsAt": "2026-08-12T04:00:00Z",
            },
        ).json()
        invalid_roster = client.put(
            f"/api/sessions/{session['id']}/members",
            headers=headers,
            json={"identityIds": [candidate["id"], blocked["id"]]},
        )
        assert invalid_roster.status_code == 400
        assert client.put(
            f"/api/sessions/{session['id']}/members",
            headers=headers,
            json={"identityIds": [candidate["id"]]},
        ).status_code == 200
        assert client.post(f"/api/sessions/{session['id']}/start", headers=headers).status_code == 200

        unknown_bytes = jpeg_bytes()
        with SessionLocal() as db:
            camera = Camera(name="Gate A", rtsp_url_enc=b"pending", enabled=False, session_id=session["id"])
            db.add(camera)
            db.flush()
            camera.rtsp_url_enc = crypto.encrypt_text(
                "rtsp://127.0.0.1:1/stream", context=f"camera:{camera.id}".encode()
            )
            unknown_key = blobs.put(unknown_bytes, namespace="event")
            unknown = RecognitionEvent(
                event_type=EventType.UNKNOWN.value,
                camera_id=camera.id,
                session_id=session["id"],
                captured_at=datetime.now(timezone.utc),
                snapshot_blob_key=unknown_key,
            )
            blacklist_key = blobs.put(jpeg_bytes("#754b43"), namespace="event")
            blacklist = RecognitionEvent(
                event_type=EventType.BLACKLIST.value,
                camera_id=camera.id,
                session_id=session["id"],
                identity_id=blocked["id"],
                identity_name_snapshot="Blocked Person",
                identity_category_snapshot="blacklist",
                confidence=0.81,
                captured_at=datetime.now(timezone.utc) + timedelta(seconds=1),
                snapshot_blob_key=blacklist_key,
            )
            db.add_all([unknown, blacklist])
            db.commit()
            unknown_id = unknown.id

        snapshot = client.get(f"/api/events/{unknown_id}/snapshot")
        assert snapshot.status_code == 200
        assert snapshot.content == unknown_bytes
        resolved = client.post(
            f"/api/events/{unknown_id}/resolve",
            headers=headers,
            json={"identityId": candidate["id"], "note": "Passport checked at desk"},
        )
        assert resolved.status_code == 200
        assert resolved.json()["resolvedIdentityName"] == "Jane Candidate"
        assert resolved.json()["resolutionNote"].startswith("Manually matched by System Administrator")
        with SessionLocal() as db:
            attendance = db.query(Attendance).filter_by(
                session_id=session["id"], identity_id=candidate["id"]
            ).one()
            assert attendance.source == "manual"
            assert "Passport checked at desk" in (attendance.remarks or "")

        report = client.get(f"/api/reports/sessions/{session['id']}.xlsx")
        assert report.status_code == 200
        workbook = load_workbook(io.BytesIO(report.content))
        assert workbook.sheetnames == ["Summary", "Attendance", "Unknowns", "Blacklist"]
        assert workbook["Attendance"]["C2"].value == "Present"
        assert workbook["Attendance"]["D2"].value == "Manual"
        assert workbook["Unknowns"]["C2"].value == "Unknown"
        assert workbook["Blacklist"]["C2"].value == "Blacklist"

        with client.websocket_connect("/api/ws/events") as websocket:
            assert websocket.receive_json()["type"] == "connected"
            event_hub.publish("camera.health", {"cameraId": 55, "connected": True})
            message = websocket.receive_json()
            assert message["type"] == "camera.health"
            assert message["payload"]["cameraId"] == 55

        deleted = client.post(
            "/api/events/bulk-delete",
            headers=headers,
            json={"eventTypes": ["unknown"], "confirmation": "DELETE"},
        )
        assert deleted.status_code == 200
        assert deleted.json() == {"deleted": 1}
        assert client.get(f"/api/events/{unknown_id}/snapshot").status_code == 404

        operator = client.post(
            "/api/users",
            headers=headers,
            json={
                "username": "operator1",
                "displayName": "Hall Operator",
                "password": "OperatorPassword123!",
                "role": "operator",
            },
        )
        assert operator.status_code == 201
        assert client.post("/api/auth/logout", headers=headers).status_code == 204
        login = client.post(
            "/api/auth/login",
            json={"username": "operator1", "password": "OperatorPassword123!"},
        )
        operator_headers = {"X-CSRF-Token": login.json()["csrfToken"]}
        assert client.post(
            "/api/identities",
            headers=operator_headers,
            json={"displayName": "Forbidden", "category": "whitelist"},
        ).status_code == 403
        assert client.get(f"/api/reports/sessions/{session['id']}.xlsx").status_code == 200


def test_blank_upload_is_reviewed_without_persisting_original() -> None:
    with TestClient(app) as client:
        login = client.post(
            "/api/auth/login",
            json={"username": "admin", "password": "StrongPassword123!"},
        )
        csrf = login.json()["csrfToken"]
        image = io.BytesIO()
        Image.new("RGB", (300, 300), "white").save(image, "PNG")
        response = client.post(
            "/api/enrollments/analyze",
            headers={"X-CSRF-Token": csrf},
            files={"files": ("no_face.png", image.getvalue(), "image/png")},
        )
        assert response.status_code == 200
        item = response.json()["items"][0]
        assert item["suggestedName"] == "no face"
        assert item["faces"] == []
        assert item["message"] == "No face detected"
